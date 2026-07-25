"""
XLSX parser — extracts data from Excel spreadsheets using openpyxl.

Each sheet is converted to a markdown table with the sheet name as a heading.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.core.constants import SourceType
from app.core.exceptions import ValidationError
from app.core.logging import logger
from app.services.parsers import BaseParser


def _sheet_to_markdown(
    sheet_name: str,
    rows: list[list[str]],
) -> str:
    """Convert a list of rows into a markdown section with a heading and table.

    Args:
        sheet_name: Name of the worksheet (used as the heading).
        rows: List of rows, where each row is a list of cell value strings.

    Returns:
        Markdown string with heading and table, or empty string if no data.
    """
    if not rows:
        return ""

    # Determine column count from the widest row
    col_count = max(len(row) for row in rows)

    # Normalise all rows to the same column count
    normalised: list[list[str]] = []
    for row in rows:
        padded = row + [""] * (col_count - len(row))
        normalised.append(padded)

    lines: list[str] = []

    # Sheet name as heading
    lines.append(f"## {sheet_name}")
    lines.append("")

    # Header row
    header = normalised[0]
    lines.append("| " + " | ".join(header) + " |")

    # Separator row
    lines.append("| " + " | ".join("---" for _ in header) + " |")

    # Data rows
    for row in normalised[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)


class XlsxParser(BaseParser):
    """Extract data from XLSX spreadsheets using openpyxl."""

    supported_types: set[SourceType] = {SourceType.XLSX}

    async def parse(self, data: bytes, *, filename: str = "") -> str:
        """Parse raw XLSX bytes and return markdown tables.

        Each worksheet is rendered as a separate markdown section with a
        ``## Sheet Name`` heading followed by the table content.

        Args:
            data: Raw XLSX file content.
            filename: Original filename (used for logging).

        Returns:
            Markdown-formatted text with one section per worksheet.

        Raises:
            ValidationError: If the file cannot be parsed.
        """
        logger.info(
            "xlsx_parse_start",
            filename=filename,
            file_size_bytes=len(data),
        )

        try:
            wb = load_workbook(
                filename=BytesIO(data),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise ValidationError(
                f"Failed to open XLSX file '{filename}': {exc}"
            ) from exc

        sections: list[str] = []
        total_rows = 0

        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows: list[list[str]] = []

                for row in ws.iter_rows():
                    cell_values = [
                        str(cell.value).strip().replace("|", "\\|")
                        if cell.value is not None
                        else ""
                        for cell in row
                    ]
                    # Skip entirely empty rows
                    if any(v for v in cell_values):
                        rows.append(cell_values)

                total_rows += len(rows)

                md_section = _sheet_to_markdown(sheet_name, rows)
                if md_section:
                    sections.append(md_section)
        finally:
            wb.close()

        result = "\n\n".join(sections)

        logger.info(
            "xlsx_parse_complete",
            filename=filename,
            sheet_count=len(sections),
            total_rows=total_rows,
            char_count=len(result),
        )

        return result
