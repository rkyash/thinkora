"""
PowerPoint (.pptx) parser — extracts slide text and speaker notes.

Uses the ``python-pptx`` library (imported as ``pptx``) to iterate over
slides, collect shape text and speaker notes, and format everything as
structured markdown.
"""

from __future__ import annotations

import io

from pptx import Presentation

from app.core.constants import SourceType
from app.core.exceptions import ValidationError
from app.core.logging import logger
from app.services.parsers import BaseParser


class PptxParser(BaseParser):
    """Extract text and speaker notes from PowerPoint presentations."""

    supported_types: set[SourceType] = {SourceType.PPTX}

    async def parse(self, data: bytes, *, filename: str = "") -> str:
        """Parse a PPTX file and return markdown-formatted slide content.

        Each slide is rendered as::

            ## Slide N
            <shape text>

            **Notes:** <speaker notes>

        Args:
            data: Raw .pptx file content.
            filename: Original filename (used for logging).

        Returns:
            Markdown string with all slide text and notes.

        Raises:
            ValidationError: If the PPTX file is corrupt or unreadable.
        """
        logger.info("pptx_parse_start", filename=filename, size=len(data))

        try:
            prs = Presentation(io.BytesIO(data))
        except Exception as exc:
            logger.error("pptx_open_failed", filename=filename, error=str(exc))
            raise ValidationError(
                f"Failed to open PPTX file '{filename}': {exc}"
            ) from exc

        sections: list[str] = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            # --- Collect shape text -----------------------------------------
            text_parts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        para_text = paragraph.text.strip()
                        if para_text:
                            text_parts.append(para_text)

            slide_text = "\n".join(text_parts)

            # --- Collect speaker notes --------------------------------------
            notes_text = ""
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                notes_text = notes_frame.text.strip()

            # --- Format slide section ---------------------------------------
            section = f"## Slide {slide_num}\n{slide_text}"
            if notes_text:
                section += f"\n\n**Notes:** {notes_text}"

            sections.append(section)

        result = "\n\n".join(sections)

        logger.info(
            "pptx_parse_complete",
            filename=filename,
            slides=len(prs.slides),
            length=len(result),
        )
        return result
