"""
DOCX parser — extracts text from Word documents using python-docx.

Preserves heading structure and formats tables as markdown tables.
"""

from __future__ import annotations

import io

import docx

from app.core.constants import SourceType
from app.core.exceptions import ValidationError
from app.core.logging import logger
from app.services.parsers import BaseParser

# Mapping from python-docx heading style names to markdown heading levels.
_HEADING_LEVEL_MAP: dict[str, int] = {
    "Heading 1": 1,
    "Heading 2": 2,
    "Heading 3": 3,
    "Heading 4": 4,
    "Heading 5": 5,
    "Heading 6": 6,
    "Title": 1,
    "Subtitle": 2,
}


def _table_to_markdown(table: docx.table.Table) -> str:
    """Convert a python-docx Table to a markdown table string."""
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [cell.text.strip().replace("|", "\\|") for cell in row.cells]
        rows.append(cells)

    if not rows:
        return ""

    # Build markdown table
    lines: list[str] = []

    # Header row
    header = rows[0]
    lines.append("| " + " | ".join(header) + " |")

    # Separator row
    lines.append("| " + " | ".join("---" for _ in header) + " |")

    # Data rows
    for row in rows[1:]:
        # Pad or trim row to match header column count
        while len(row) < len(header):
            row.append("")
        lines.append("| " + " | ".join(row[: len(header)]) + " |")

    return "\n".join(lines)


class DocxParser(BaseParser):
    """Extract text from DOCX documents using python-docx."""

    supported_types: set[SourceType] = {SourceType.DOCX}

    async def parse(self, data: bytes, *, filename: str = "") -> str:
        """Parse raw DOCX bytes and return extracted text.

        Paragraphs are extracted in order, headings become markdown headings,
        and tables are converted to markdown table format.

        Args:
            data: Raw DOCX file content.
            filename: Original filename (used for logging).

        Returns:
            Extracted text with headings and markdown tables.

        Raises:
            ValidationError: If the file cannot be parsed.
        """
        logger.info(
            "docx_parse_start",
            filename=filename,
            file_size_bytes=len(data),
        )

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise ValidationError(
                f"Failed to open DOCX file '{filename}': {exc}"
            ) from exc

        # Build an ordered list of blocks (paragraphs and tables) as they
        # appear in the document body.  python-docx exposes the XML element
        # tree, so we walk through the body children and dispatch accordingly.
        sections: list[str] = []
        paragraph_count = 0
        table_count = 0

        # Map XML elements to their corresponding python-docx objects for
        # O(1) lookup while iterating the body children.
        para_map = {p._element: p for p in document.paragraphs}
        table_map = {t._element: t for t in document.tables}

        for child in document.element.body:
            if child in para_map:
                paragraph = para_map[child]
                text = paragraph.text.strip()
                if not text:
                    continue

                # Detect heading style
                style_name = paragraph.style.name if paragraph.style else ""
                level = _HEADING_LEVEL_MAP.get(style_name, 0)

                if level:
                    sections.append(f"{'#' * level} {text}")
                else:
                    sections.append(text)
                paragraph_count += 1

            elif child in table_map:
                table = table_map[child]
                md_table = _table_to_markdown(table)
                if md_table:
                    sections.append(md_table)
                    table_count += 1

        result = "\n\n".join(sections)

        logger.info(
            "docx_parse_complete",
            filename=filename,
            paragraph_count=paragraph_count,
            table_count=table_count,
            char_count=len(result),
        )

        return result
